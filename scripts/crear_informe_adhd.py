from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "documentos" / "resultados_random_forest_adhd.docx"
GRAFICOS = RAIZ / "resultados_adhd"


def fuente(run, tamanio=11, negrita=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(tamanio)
    run.bold = negrita
    if color:
        run.font.color.rgb = RGBColor(*color)


def sombrear(celda, color):
    propiedades = celda._tc.get_or_add_tcPr()
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:fill"), color)
    propiedades.append(sombreado)


def bordes(celda, color="D9E2F3"):
    propiedades = celda._tc.get_or_add_tcPr()
    elementos = propiedades.first_child_found_in("w:tcBorders")
    if elementos is None:
        elementos = OxmlElement("w:tcBorders")
        propiedades.append(elementos)
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), "6")
        borde.set(qn("w:color"), color)
        elementos.append(borde)


def parrafo(documento, texto, negrita_inicio=None):
    p = documento.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if negrita_inicio and texto.startswith(negrita_inicio):
        r = p.add_run(negrita_inicio)
        fuente(r, negrita=True)
        r = p.add_run(texto[len(negrita_inicio):])
        fuente(r)
    else:
        r = p.add_run(texto)
        fuente(r)
    return p


def titulo(documento, texto, nivel=1):
    p = documento.add_paragraph(style=f"Heading {nivel}")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(texto)
    fuente(r, 14 if nivel == 1 else 12, negrita=True, color=(0, 0, 0))
    return p


def figura(documento, archivo, leyenda):
    ruta = GRAFICOS / archivo
    if not ruta.exists():
        raise FileNotFoundError(f"No se encontró el gráfico: {ruta}")
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(ruta), width=Cm(15.2))
    caption = documento.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    r = caption.add_run(leyenda)
    fuente(r, 10)
    r.italic = True


def tabla_metricas(documento):
    filas = [
        ("Accuracy", "0.8070", "Proporción total de clasificaciones correctas."),
        ("Balanced accuracy", "0.8064", "Promedio equilibrado del acierto en ambas clases."),
        ("Precision", "0.7894", "Proporción de positivos predichos que fueron positivos."),
        ("Recall o sensibilidad", "0.7978", "Proporción de positivos identificados correctamente."),
        ("Especificidad", "0.8150", "Proporción de negativos identificados correctamente."),
        ("F1 score", "0.7936", "Equilibrio entre precisión y sensibilidad."),
        ("ROC AUC", "0.8594", "Capacidad de separar las clases positiva y negativa."),
        ("PR AUC", "0.8267", "Desempeño de precisión y sensibilidad en distintos umbrales."),
    ]
    tabla = documento.add_table(rows=1, cols=3)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.style = "Table Grid"
    encabezados = ["Métrica", "Resultado", "Interpretación"]
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        sombrear(celda, "1F4E78")
        celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        fuente(r, 10, negrita=True, color=(255, 255, 255))
        bordes(celda)
    for indice, fila in enumerate(filas):
        celdas = tabla.add_row().cells
        for posicion, texto in enumerate(fila):
            if indice % 2 == 1:
                sombrear(celdas[posicion], "F3F7FB")
            celdas[posicion].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = celdas[posicion].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if posicion == 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(texto)
            fuente(r, 9)
            bordes(celdas[posicion])
    return tabla


def main():
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    documento = Document()
    seccion = documento.sections[0]
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(2.7)
    seccion.right_margin = Cm(2.7)

    estilos = documento.styles
    estilos["Normal"].font.name = "Calibri"
    estilos["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    estilos["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    estilos["Normal"].font.size = Pt(11)
    for nombre in ("Title", "Heading 1", "Heading 2"):
        estilos[nombre].font.color.rgb = RGBColor(0, 0, 0)

    portada = documento.add_paragraph(style="Title")
    portada.alignment = WD_ALIGN_PARAGRAPH.CENTER
    portada.paragraph_format.space_before = Pt(130)
    r = portada.add_run("Resultados del modelo Random Forest para ADHD Rating Scale IV")
    fuente(r, 20, negrita=True)
    subtitulo = documento.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitulo.paragraph_format.space_before = Pt(12)
    r = subtitulo.add_run("Versión para profesores")
    fuente(r, 13)
    documento.add_page_break()

    titulo(documento, "Resultados del modelo Random Forest")
    parrafo(documento, "En esta sección se presentan los resultados obtenidos mediante el algoritmo Random Forest Classifier aplicado a la ADHD Rating Scale IV, versión para profesores. El objetivo de esta primera implementación fue comprobar el funcionamiento técnico del proceso de generación de datos, entrenamiento, evaluación y visualización de un modelo de clasificación.")
    parrafo(documento, "Para esta prueba se generaron 10 000 registros sintéticos. Cada registro contiene las respuestas de los 18 ítems de la escala, con valores entre 0 y 3, donde 0 representa nunca o rara vez y 3 representa una frecuencia alta de la conducta evaluada. La variable objetivo fue construida como un indicador sintético de tamizaje positivo: se consideró positivo cuando existían seis o más ítems con puntuación mayor o igual a 2 en la dimensión de inatención o en la dimensión de hiperactividad e impulsividad. Se incorporó un 12 por ciento de ruido en las etiquetas para evitar una clasificación artificialmente perfecta.")
    parrafo(documento, "Los datos fueron divididos en 80 por ciento para entrenamiento, equivalente a 8 000 registros, y 20 por ciento para pruebas, equivalente a 2 000 registros. El modelo se configuró con 500 árboles de decisión, ponderación balanceada de clases y una semilla aleatoria fija de 42, con el propósito de garantizar resultados reproducibles.")

    titulo(documento, "Métricas de evaluación", nivel=2)
    parrafo(documento, "Como se observa en la Figura 1, el modelo obtuvo una exactitud de 80.70 por ciento, lo que indica que clasificó correctamente aproximadamente ocho de cada diez registros del conjunto de prueba. La precisión fue de 78.94 por ciento, mientras que el recall o sensibilidad alcanzó 79.78 por ciento. El F1 score fue de 79.36 por ciento y el valor de ROC AUC fue de 85.94 por ciento.")
    tabla_metricas(documento)
    figura(documento, "01_metricas.png", "Figura 1. Métricas de evaluación del modelo Random Forest.")

    titulo(documento, "Matriz de confusión", nivel=2)
    parrafo(documento, "La Figura 2 presenta la matriz de confusión del modelo. Se identificaron correctamente 872 casos negativos verdaderos y 742 casos positivos verdaderos. Asimismo, se produjeron 198 falsos positivos, es decir, registros clasificados como positivos cuando su etiqueta sintética era negativa. Se registraron 188 falsos negativos, correspondientes a casos positivos que el modelo no identificó.")
    parrafo(documento, "La especificidad obtenida fue de 81.50 por ciento. Aunque el modelo presenta un desempeño adecuado sobre estos datos simulados, los falsos negativos deben ser considerados con atención cuando se realice una evaluación con datos reales.")
    figura(documento, "02_matriz_confusion.png", "Figura 2. Matriz de confusión del modelo Random Forest.")

    titulo(documento, "Curvas ROC y Precision Recall", nivel=2)
    parrafo(documento, "Como se observa en la Figura 3, la curva ROC se ubica por encima de la línea diagonal que representa una clasificación aleatoria. El AUC de 0.859 demuestra que el modelo posee capacidad de discriminación entre ambas clases en el conjunto de prueba.")
    figura(documento, "03_curva_roc.png", "Figura 3. Curva ROC del modelo Random Forest.")
    parrafo(documento, "La Figura 4 corresponde a la curva Precision Recall, cuyo valor promedio fue de 0.827. Esta gráfica permite evaluar la relación entre la precisión y la sensibilidad para distintos umbrales de decisión. Los resultados muestran que el modelo mantiene un desempeño consistente al identificar la clase positiva sintética.")
    figura(documento, "04_curva_precision_recall.png", "Figura 4. Curva Precision Recall del modelo Random Forest.")

    titulo(documento, "Calibración del modelo", nivel=2)
    parrafo(documento, "La Figura 5 muestra la curva de calibración, la cual compara la probabilidad predicha por el modelo con la proporción real de casos positivos. Se observan variaciones respecto de la línea de calibración perfecta, principalmente en probabilidades intermedias. Por ello, las probabilidades estimadas deben interpretarse con precaución y requerirán calibración adicional cuando se disponga de datos reales.")
    figura(documento, "05_calibracion.png", "Figura 5. Curva de calibración del modelo Random Forest.")

    titulo(documento, "Importancia de las características", nivel=2)
    parrafo(documento, "La Figura 6 presenta la importancia relativa de los 18 ítems de la escala. Los ítems con mayor importancia fueron el ítem 1, relacionado con la atención a detalles; el ítem 14, referido a responder impulsivamente; y el ítem 7, relacionado con seguir instrucciones. Sin embargo, las diferencias entre las importancias son reducidas, debido a que los datos usados fueron generados aleatoriamente y todos los ítems participan de manera similar en la etiqueta sintética.")
    figura(documento, "06_importancia_items.png", "Figura 6. Importancia de las características de ADHD Rating Scale IV.")

    titulo(documento, "Validación cruzada", nivel=2)
    parrafo(documento, "La Figura 7 presenta una validación cruzada estratificada de cinco particiones. Los resultados se mantuvieron relativamente estables entre los distintos folds: accuracy promedio de 78.25 por ciento, precisión de 77.50 por ciento, sensibilidad de 75.03 por ciento, F1 score de 76.24 por ciento y ROC AUC de 83.90 por ciento. Esta estabilidad indica que el modelo no depende exclusivamente de una única división de entrenamiento y prueba.")
    figura(documento, "07_validacion_cruzada.png", "Figura 7. Validación cruzada estratificada del modelo Random Forest.")

    titulo(documento, "Limitación de la prueba", nivel=2)
    parrafo(documento, "Los datos y las etiquetas empleados en esta fase fueron simulados. En consecuencia, los resultados validan el funcionamiento técnico del flujo de machine learning, pero no constituyen evidencia de desempeño clínico ni deben utilizarse como diagnóstico. La siguiente etapa debe incorporar respuestas reales de docentes y una etiqueta clínica o profesional validada para entrenar y evaluar el modelo como apoyo al tamizaje.")

    documento.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    main()
