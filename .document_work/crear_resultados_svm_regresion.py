from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = Path(r"C:\Users\PERCY\Desktop\resultados_random_forest_adhd.docx")
OUTPUT = ROOT / "resultados_svm_regresion_lineal_adhd.docx"
RESULTS = ROOT / "resultados_adhd"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    tc_pr.append(shade)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def add_text(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    return p


def add_figure(doc, image, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image), width=Inches(5.9))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)


def add_metrics_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.55), Inches(1.15), Inches(3.65)]
    headers = ["Métrica", "Resultado", "Interpretación"]
    for cell, width, text in zip(table.rows[0].cells, widths, headers):
        cell.width = width; cell.text = text; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78"); set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for index, (metric, value, interpretation) in enumerate(rows):
        cells = table.add_row().cells
        for cell, width, text in zip(cells, widths, [metric, value, interpretation]):
            cell.width = width; cell.text = text; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if index % 2 == 1:
                set_cell_shading(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs: run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_model_section(doc, heading, folder, values, matrix_text, roc_text, pr_text, calibration_text, importance_text, cv_text, figure_start):
    add_text(doc, heading, "Heading 1")
    add_text(doc, "Esta sección presenta los resultados del modelo " + heading +
             " aplicado al tamizaje ADHD Rating Scale IV para profesores. La evaluación usa la misma división estratificada: 24 registros para entrenamiento y 6 para prueba.")
    add_text(doc, "Métricas de evaluación", "Heading 2")
    add_metrics_table(doc, values)
    add_text(doc, f"La Figura {figure_start} resume las métricas del modelo. Los resultados deben interpretarse como una comprobación técnica del flujo de entrenamiento y no como evidencia clínica.")
    add_figure(doc, RESULTS / folder / "01_metricas.png", f"Figura {figure_start}. Métricas de evaluación del modelo {heading}.")

    add_text(doc, "Matriz de confusión", "Heading 2")
    add_text(doc, matrix_text)
    add_figure(doc, RESULTS / folder / "02_matriz_confusion.png", f"Figura {figure_start + 1}. Matriz de confusión del modelo {heading}.")

    add_text(doc, "Curvas ROC y Precision Recall", "Heading 2")
    add_text(doc, roc_text)
    add_figure(doc, RESULTS / folder / "03_curva_roc.png", f"Figura {figure_start + 2}. Curva ROC del modelo {heading}.")
    add_text(doc, pr_text)
    add_figure(doc, RESULTS / folder / "04_curva_precision_recall.png", f"Figura {figure_start + 3}. Curva Precision Recall del modelo {heading}.")

    add_text(doc, "Calibración del modelo", "Heading 2")
    add_text(doc, calibration_text)
    add_figure(doc, RESULTS / folder / "05_calibracion.png", f"Figura {figure_start + 4}. Curva de calibración del modelo {heading}.")

    add_text(doc, "Características del modelo", "Heading 2")
    add_text(doc, importance_text)
    add_figure(doc, RESULTS / folder / "06_importancia_items.png", f"Figura {figure_start + 5}. Características evaluadas por el modelo {heading}.")

    add_text(doc, "Validación cruzada", "Heading 2")
    add_text(doc, cv_text)
    add_figure(doc, RESULTS / folder / "07_validacion_cruzada.png", f"Figura {figure_start + 6}. Validación cruzada estratificada del modelo {heading}.")


def main():
    doc = Document(TEMPLATE)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Resultados de los modelos SVM y regresión logística para ADHD Rating Scale IV")
    subtitle = doc.add_paragraph("Versión para profesores")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)

    add_text(doc, "Resultados de modelos de clasificación", "Heading 1")
    add_text(doc, "El presente documento reúne los resultados obtenidos con Support Vector Machine SVM y regresión logística, dos algoritmos de clasificación entrenados para el flujo de tamizaje ADHD Rating Scale IV. Ambos utilizan las respuestas de los 18 ítems, con valores entre 0 y 3, y estiman la probabilidad de un tamizaje positivo.")
    add_text(doc, "El conjunto disponible contiene 30 registros sin respuestas reales. Por ello, el entrenamiento se realizó con datos sintéticos reproducibles generados por el script del proyecto. Los resultados sirven para comparar el funcionamiento técnico de los modelos; no constituyen validación clínica.")

    svm_metrics = [
        ("Accuracy", "0.6667", "Proporción total de clasificaciones correctas."),
        ("Balanced accuracy", "0.5000", "Promedio equilibrado del acierto en ambas clases."),
        ("Precision", "0.0000", "Proporción de positivos predichos que fueron positivos."),
        ("Recall o sensibilidad", "0.0000", "Proporción de positivos identificados correctamente."),
        ("Especificidad", "1.0000", "Proporción de negativos identificados correctamente."),
        ("F1 score", "0.0000", "Equilibrio entre precisión y sensibilidad."),
        ("ROC AUC", "0.6250", "Capacidad de separar las clases positiva y negativa."),
        ("PR AUC", "0.5000", "Desempeño de precisión y sensibilidad en distintos umbrales."),
    ]
    add_model_section(
        doc, "SVM", "svm", svm_metrics,
        "La Figura 2 muestra cuatro negativos verdaderos y dos falsos negativos. No se produjeron falsos positivos ni positivos verdaderos. Aunque la especificidad fue de 100 por ciento, el modelo no identificó los casos positivos del conjunto de prueba con el umbral utilizado.",
        "La Figura 3 muestra una curva ROC con AUC de 0.625. El resultado se encuentra por encima del azar, pero la capacidad de discriminación es moderada y debe confirmarse con datos reales.",
        "La Figura 4 presenta una curva Precision Recall con AP de 0.500. El valor refleja una capacidad limitada para mantener precisión al recuperar casos positivos en esta prueba pequeña.",
        "La Figura 5 compara probabilidades predichas y proporciones reales. Con solo seis casos de prueba, la curva debe considerarse exploratoria y no una calibración definitiva.",
        "La Figura 6 indica que la SVM con kernel RBF no entrega coeficientes lineales por ítem. El modelo utiliza una frontera no lineal construida a partir de todos los ítems escalados.",
        "La Figura 7 resume la validación cruzada estratificada de cinco particiones. El accuracy promedio fue de 58.00 por ciento, el recall de 20.00 por ciento y el F1 promedio de 8.00 por ciento, con variación entre particiones debido al tamaño reducido de la muestra.",
        1,
    )

    doc.add_page_break()
    logistic_metrics = [
        ("Accuracy", "0.6667", "Proporción total de clasificaciones correctas."),
        ("Balanced accuracy", "0.6250", "Promedio equilibrado del acierto en ambas clases."),
        ("Precision", "0.5000", "Proporción de positivos predichos que fueron positivos."),
        ("Recall o sensibilidad", "0.5000", "Proporción de positivos identificados correctamente."),
        ("Especificidad", "0.7500", "Proporción de negativos identificados correctamente."),
        ("F1 score", "0.5000", "Equilibrio entre precisión y sensibilidad."),
        ("ROC AUC", "0.6250", "Capacidad de separar las clases positiva y negativa."),
        ("PR AUC", "0.7000", "Desempeño de precisión y sensibilidad en distintos umbrales."),
    ]
    add_model_section(
        doc, "Regresión logística", "regresion_lineal", logistic_metrics,
        "La Figura 9 presenta tres negativos verdaderos, un falso positivo, un falso negativo y un positivo verdadero. A diferencia de la SVM, este modelo logró identificar un caso positivo del conjunto de prueba.",
        "La Figura 10 muestra una curva ROC con AUC de 0.625. La discriminación es moderada en esta muestra y requiere evaluación posterior con registros reales.",
        "La Figura 11 presenta una curva Precision Recall con AP de 0.700. Este resultado fue superior al de SVM para la recuperación de la clase positiva en la misma partición de prueba.",
        "La Figura 12 compara las probabilidades estimadas con las etiquetas observadas. La cantidad reducida de casos limita cualquier conclusión sobre la calibración final del modelo.",
        "La Figura 13 muestra el valor absoluto de los coeficientes del clasificador lineal. Estos valores describen la influencia relativa de cada ítem dentro de este modelo sintético y no deben interpretarse como relevancia clínica.",
        "La Figura 14 presenta la validación cruzada estratificada de cinco particiones. El accuracy promedio fue de 37.00 por ciento, el recall de 20.00 por ciento y el F1 promedio de 6.67 por ciento. La variabilidad confirma la necesidad de ampliar el conjunto de datos.",
        8,
    )

    add_text(doc, "Comparación y limitación de la prueba", "Heading 1")
    add_text(doc, "En la partición de prueba, ambos modelos alcanzaron 66.67 por ciento de accuracy. La regresión logística mostró un F1 de 0.500 y detectó un caso positivo, mientras que SVM tuvo F1 de 0.000 y no detectó positivos con el umbral actual. Por esta razón, la regresión logística presenta un comportamiento más útil para esta comparación técnica inicial.")
    add_text(doc, "Los datos y etiquetas son sintéticos porque el archivo fuente aún no contiene respuestas reales. En consecuencia, las métricas solo validan la ejecución del flujo de machine learning. Antes de utilizar cualquiera de los modelos como apoyo al tamizaje, se deben incorporar respuestas reales de docentes y etiquetas profesionales validadas, además de realizar una evaluación con una muestra más amplia.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
